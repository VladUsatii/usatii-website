from __future__ import annotations

import re
from pathlib import Path

import pdfplumber
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_RIGHT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas as pdfcanvas
from reportlab.platypus import BaseDocTemplate, Frame, Image, PageBreak, PageTemplate, Paragraph, Spacer, Table, TableStyle

from build_usatii_proposal import configure_doc, configure_section, force_inter


ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path("/Users/vladusatii/Downloads/Proposal _ Jul 30 2026.pdf")
DOCX_OUT = ROOT / "tmp" / "pdfs" / "USATII_MEDIA_Windmill_Reddit_Intelligence_Proposal.docx"
PDF_OUT = ROOT / "output" / "pdf" / "USATII_MEDIA_Windmill_Reddit_Intelligence_Proposal.pdf"
INTER = Path("/Users/vladusatii/Library/Fonts/Inter-Regular.ttf")
INTER_BOLD = Path("/Users/vladusatii/Library/Fonts/Inter-Bold.ttf")
LOGO = ROOT / "assets" / "usatii-media-header-logo.png"

SECTION_HEADINGS = {
    "Workflow", "Severity levels", "Knowledge base", "Step 3: Standard founder queue.",
    "Founder-selection criteria", "Initial routing method", "Case information", "Reddit AMA",
    "Discord and Slack communities", "Weekly operating report", "Final pilot report", "Schedule",
    "Deliverables", "Unified conversation queue. Accepted when:",
    "Classification system. Accepted when:", "Approved knowledge base. Accepted when:",
    "Response drafting workflow. Accepted when:", "Founder response queue. Accepted when:",
    "Support and product routing. Accepted when:", "Founder content. Accepted when:",
    "AMA. Accepted when:", "Reporting. Accepted when:", "Client responsibilities. Windmill will:",
    "USATII MEDIA will:", "Commercial Terms", "Payment schedule:",
}


def clean(text: str) -> str:
    # The source PDF uses private-use glyphs for punctuation in a subset font.
    return (text.replace("\ue092", ":").replace("\ue028", "-")
            .replace("\u00ad", "").replace("\uf0b7", "●").strip())


def extracted_lines() -> list[str]:
    lines: list[str] = []
    with pdfplumber.open(SOURCE) as pdf:
        for page in pdf.pages:
            text = page.extract_text(x_tolerance=2, y_tolerance=2) or ""
            lines.extend(clean(line) for line in text.splitlines())
    return [line for line in lines if line]


def logical_items(lines: list[str]) -> list[tuple[str, str]]:
    items: list[tuple[str, str]] = []
    current = ""

    def flush() -> None:
        nonlocal current
        if current:
            items.append(("paragraph", current.strip()))
            current = ""

    for line in lines:
        if line == "—":
            flush()
            continue
        if line in SECTION_HEADINGS or re.fullmatch(r"Days \d+[-:]\d+:?", line):
            flush()
            items.append(("heading", line))
        elif line.startswith(("● ", "○ ")):
            flush()
            items.append(("bullet2" if line.startswith("○ ") else "bullet", line[2:].strip()))
        elif re.match(r"^\d+\.\s", line):
            flush()
            items.append(("number", line))
        else:
            current = f"{current} {line}".strip()
    flush()
    return items


def add_styled_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    match = re.match(r"^([^.!?]{2,55}[.:])\s+(.*)$", text)
    if match and match.group(1) not in {"USATII MEDIA:", "Windmill:"}:
        r = p.add_run(match.group(1) + " ")
        r.bold = True
        p.add_run(match.group(2))
    else:
        p.add_run(text)


def build() -> Path:
    lines = extracted_lines()
    title = lines[0]
    metadata = lines[1:7]
    body = lines[7:]

    doc = Document()
    configure_doc(
        doc,
        first_page=True,
        document_title="Reddit Intelligence Proposal",
        date_value="July 30, 2026",
        prepared_for="Windmill · Michael Mayer",
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run(title.upper())
    r.bold = True
    r.font.size = Pt(17)

    for line in metadata:
        add_styled_paragraph(doc, line)

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section, first_page=False, document_title="Reddit Intelligence and Founder Content Pilot")

    for kind, text in logical_items(body):
        if kind == "heading":
            p = doc.add_paragraph(style="Heading 1" if text in {"Workflow", "Schedule", "Deliverables", "Commercial Terms"} else "Heading 2")
            p.add_run(text)
        elif kind.startswith("bullet"):
            p = doc.add_paragraph(style="List Bullet 2" if kind == "bullet2" else "List Bullet")
            p.paragraph_format.space_after = Pt(2)
            p.add_run(text)
        elif kind == "number":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            p.paragraph_format.first_line_indent = Pt(-18)
            p.paragraph_format.space_after = Pt(2)
            p.add_run(text)
        else:
            add_styled_paragraph(doc, text)

    force_inter(doc)
    DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_OUT)
    return DOCX_OUT


def build_pdf() -> Path:
    pdfmetrics.registerFont(TTFont("Inter", str(INTER)))
    pdfmetrics.registerFont(TTFont("Inter-Bold", str(INTER_BOLD)))
    lines = extracted_lines()
    title, metadata, body = lines[0], lines[1:7], lines[7:]
    PDF_OUT.parent.mkdir(parents=True, exist_ok=True)

    doc = BaseDocTemplate(str(PDF_OUT), pagesize=letter, leftMargin=.60*inch, rightMargin=.60*inch)
    first = Frame(.60*inch, .56*inch, 7.30*inch, 8.32*inch, id="first")
    normal = Frame(.60*inch, .56*inch, 7.30*inch, 9.65*inch, id="normal")

    def header_footer(canvas, d):
        canvas.saveState()
        canvas.setFont("Inter", 7)
        canvas.setFillColor(colors.HexColor("#62666D"))
        canvas.drawRightString(7.90*inch, .28*inch, f"Page {d.page}")
        if d.page == 1:
            canvas.drawImage(str(LOGO), .60*inch, 9.92*inch, width=1.55*inch, height=.28*inch, preserveAspectRatio=True, anchor="sw", mask="auto")
            canvas.setFont("Inter", 8.5); canvas.setFillColor(colors.HexColor("#111318"))
            for i, text in enumerate(("Usatii Media", "usatii.com", "vlad@usatii.com")):
                canvas.drawString(.60*inch, (9.80-i*.16)*inch, text)
            x, y, w, h = 4.22*inch, 9.48*inch, 3.68*inch, .80*inch
            data = [["Date", "Document details"], ["July 30, 2026", "Reddit Intelligence Proposal"],
                    ["Prepared by", "Prepared for"], ["Vladislav Usatii", "Windmill · Michael Mayer"]]
            t = Table(data, colWidths=[1.25*inch, 2.43*inch], rowHeights=[.20*inch]*4)
            t.setStyle(TableStyle([("FONTNAME",(0,0),(-1,-1),"Inter"),("FONTNAME",(0,0),(-1,0),"Inter-Bold"),
                                   ("FONTNAME",(0,2),(-1,2),"Inter-Bold"),("FONTSIZE",(0,0),(-1,-1),7),
                                   ("GRID",(0,0),(-1,-1),.4,colors.HexColor("#333333")),("VALIGN",(0,0),(-1,-1),"MIDDLE"),
                                   ("ALIGN",(0,0),(-1,0),"CENTER"),("ALIGN",(0,2),(-1,2),"CENTER"),
                                   ("LEFTPADDING",(0,0),(-1,-1),5),("RIGHTPADDING",(0,0),(-1,-1),5)]))
            t.wrapOn(canvas,w,h); t.drawOn(canvas,x,y)
        else:
            canvas.drawImage(str(LOGO), .60*inch, 10.42*inch, width=.78*inch, height=.18*inch, preserveAspectRatio=True, anchor="sw", mask="auto")
            canvas.setFont("Inter-Bold", 7); canvas.setFillColor(colors.HexColor("#111318"))
            canvas.drawRightString(7.90*inch, 10.46*inch, "Reddit Intelligence and Founder Content Pilot")
            canvas.setStrokeColor(colors.HexColor("#7C3AED")); canvas.setLineWidth(.5)
            canvas.line(.60*inch, 10.36*inch, 7.90*inch, 10.36*inch)
        canvas.restoreState()

    doc.addPageTemplates([PageTemplate("first", [first], header_footer, autoNextPageTemplate="normal"),
                          PageTemplate("normal", [normal], header_footer)])
    styles = getSampleStyleSheet()
    base = ParagraphStyle("base", parent=styles["BodyText"], fontName="Inter", fontSize=8.7, leading=10.3,
                          textColor=colors.HexColor("#111318"), spaceAfter=4)
    h1 = ParagraphStyle("h1", parent=base, fontName="Inter-Bold", fontSize=12.4, leading=14, spaceBefore=10, spaceAfter=5, keepWithNext=True)
    h2 = ParagraphStyle("h2", parent=base, fontName="Inter-Bold", fontSize=10.4, leading=12, spaceBefore=7, spaceAfter=4, keepWithNext=True)
    cover = ParagraphStyle("cover", parent=base, fontName="Inter-Bold", fontSize=17, leading=20, alignment=TA_CENTER, spaceAfter=18)
    bullet = ParagraphStyle("bullet", parent=base, leftIndent=14, firstLineIndent=-8, bulletIndent=4, spaceAfter=2)
    number = ParagraphStyle("number", parent=base, leftIndent=18, firstLineIndent=-18, spaceAfter=2)

    story = [Paragraph(title.upper(), cover), Spacer(1, 4)]
    for line in metadata:
        m = re.match(r"^([^:]+:)(.*)$", line)
        story.append(Paragraph(f"<b>{m.group(1)}</b>{m.group(2)}" if m else line, base))
    story.append(PageBreak())
    for kind, text in logical_items(body):
        safe = text.replace("&", "&amp;").replace("●", " -").replace("○", " -")
        if kind == "heading":
            story.append(Paragraph(safe, h1 if text in {"Workflow", "Schedule", "Deliverables", "Commercial Terms"} else h2))
        elif kind.startswith("bullet"):
            story.append(Paragraph(safe, bullet, bulletText="•" if kind == "bullet" else "◦"))
        elif kind == "number":
            story.append(Paragraph(safe, number))
        else:
            m = re.match(r"^([^.!?]{2,55}[.:])\s+(.*)$", safe)
            story.append(Paragraph(f"<b>{m.group(1)}</b> {m.group(2)}" if m else safe, base))
    doc.build(story)
    return PDF_OUT


def build_positioned_pdf() -> Path:
    """Re-typeset the source at its original positions using the USATII page furniture."""
    pdfmetrics.registerFont(TTFont("Inter", str(INTER)))
    pdfmetrics.registerFont(TTFont("Inter-Bold", str(INTER_BOLD)))
    PDF_OUT.parent.mkdir(parents=True, exist_ok=True)
    c = pdfcanvas.Canvas(str(PDF_OUT), pagesize=letter, pageCompression=1)

    with pdfplumber.open(SOURCE) as source:
        for page_number, page in enumerate(source.pages, 1):
            # Established USATII template header/footer.
            c.setFillColor(colors.HexColor("#62666D")); c.setFont("Inter", 7)
            c.drawRightString(7.90*inch, .28*inch, f"Page {page_number}")
            if page_number == 1:
                c.drawImage(str(LOGO), .60*inch, 9.92*inch, width=1.55*inch, height=.28*inch,
                            preserveAspectRatio=True, anchor="sw", mask="auto")
                c.setFillColor(colors.HexColor("#111318")); c.setFont("Inter", 8.5)
                for i, value in enumerate(("Usatii Media", "usatii.com", "vlad@usatii.com")):
                    c.drawString(.60*inch, (9.80-i*.16)*inch, value)
                x, y = 4.22*inch, 9.48*inch
                table = Table(
                    [["Date", "Document details"], ["July 30, 2026", "Reddit Intelligence Proposal"],
                     ["Prepared by", "Prepared for"], ["Vladislav Usatii", "Windmill · Michael Mayer"]],
                    colWidths=[1.25*inch, 2.43*inch], rowHeights=[.20*inch]*4,
                )
                table.setStyle(TableStyle([
                    ("FONTNAME",(0,0),(-1,-1),"Inter"),("FONTNAME",(0,0),(-1,0),"Inter-Bold"),
                    ("FONTNAME",(0,2),(-1,2),"Inter-Bold"),("FONTSIZE",(0,0),(-1,-1),7),
                    ("GRID",(0,0),(-1,-1),.4,colors.HexColor("#333333")),
                    ("VALIGN",(0,0),(-1,-1),"MIDDLE"),("ALIGN",(0,0),(-1,0),"CENTER"),
                    ("ALIGN",(0,2),(-1,2),"CENTER"),("LEFTPADDING",(0,0),(-1,-1),5),
                    ("RIGHTPADDING",(0,0),(-1,-1),5),
                ]))
                table.wrapOn(c, 3.68*inch, .80*inch); table.drawOn(c, x, y)
                target_top, target_bottom = 8.75*inch, .56*inch
            else:
                c.drawImage(str(LOGO), .60*inch, 10.42*inch, width=.78*inch, height=.18*inch,
                            preserveAspectRatio=True, anchor="sw", mask="auto")
                c.setFillColor(colors.HexColor("#111318")); c.setFont("Inter-Bold", 7)
                c.drawRightString(7.90*inch, 10.46*inch, "Reddit Intelligence and Founder Content Pilot")
                c.setStrokeColor(colors.HexColor("#7C3AED")); c.setLineWidth(.5)
                c.line(.60*inch, 10.36*inch, 7.90*inch, 10.36*inch)
                target_top, target_bottom = 10.18*inch, .56*inch

            words = page.extract_words(extra_attrs=["fontname", "size"], keep_blank_chars=False)
            if words:
                source_top = min(w["top"] for w in words)
                source_bottom = max(w["bottom"] for w in words)
                source_left = min(w["x0"] for w in words)
                source_right = max(w["x1"] for w in words)
                sx = (7.30*inch) / (source_right-source_left)
                sy = (target_top-target_bottom) / (source_bottom-source_top)
                scale = min(sx, sy)
                used_width = (source_right-source_left)*scale
                left = .60*inch + (7.30*inch-used_width)/2
                for word in words:
                    text = clean(word["text"])
                    if text in {"●", "○"}:
                        radius = max(1.1, float(word["size"])*scale*.18)
                        cx = left + (word["x0"]-source_left)*scale + radius
                        cy = target_top - (word["top"]-source_top)*scale - float(word["size"])*scale*.55
                        c.setFillColor(colors.HexColor("#111318"))
                        if text == "●": c.circle(cx, cy, radius, stroke=0, fill=1)
                        else: c.circle(cx, cy, radius, stroke=1, fill=0)
                        continue
                    font = "Inter-Bold" if "bold" in word["fontname"].lower() else "Inter"
                    size = max(5.8, float(word["size"])*scale)
                    x = left + (word["x0"]-source_left)*scale
                    baseline = target_top - (word["top"]-source_top)*scale - size*.78
                    c.setFillColor(colors.HexColor("#111318")); c.setFont(font, size)
                    c.drawString(x, baseline, text)
            c.showPage()
    c.save()
    return PDF_OUT


if __name__ == "__main__":
    print(build_positioned_pdf())
